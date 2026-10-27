"""Excel report for the business side.

Three sheets people actually use: a summary for the project lead, a worklist
(Klärungsliste) with empty status/owner/comment columns for the departments,
and duplicate pairs side by side so a human can decide which record survives.
"""
from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .rules import REGISTRY

SEV_ORDER = {"blocker": 0, "error": 1, "warning": 2, "info": 3}
SEV_FILL = {
    "blocker": PatternFill("solid", fgColor="F4C7C3"),
    "error": PatternFill("solid", fgColor="FCE8B2"),
    "warning": PatternFill("solid", fgColor="FFF7D6"),
    "info": PatternFill("solid", fgColor="E8F0FE"),
}
HEAD = Font(bold=True)
WIDTH_SAMPLE_ROWS = 200
MAX_WIDTH = 60


def _sheet(ws, df, widths=None):
    ws.append(list(df.columns))
    for c in ws[1]:
        c.font = HEAD
    for row in df.itertuples(index=False):
        ws.append(["" if pd.isna(v) else v for v in row])
    ws.freeze_panes = "A2"
    if len(df):
        ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(df.columns, 1):
        w = (widths or {}).get(col)
        if w is None:
            sample = [len(str(col))] + [len(str(v)) for v in df[col].head(WIDTH_SAMPLE_ROWS)]
            w = min(max(sample) + 2, MAX_WIDTH)
        ws.column_dimensions[get_column_letter(i)].width = w


def write(result, tables, out_path, title="Stammdaten-Check"):
    f = result.findings.copy()
    f["_o"] = f["severity"].map(SEV_ORDER)
    f = f.sort_values(["_o", "rule", "object", "key"]).drop(columns=["_o", "weight"])

    wb = Workbook()
    _summary(wb.active, result, f, title)
    _worklist(wb.create_sheet("Worklist"), f)
    _duplicates(wb.create_sheet("Duplicates"), f, tables)
    _rules(wb.create_sheet("Rules"))
    wb.save(out_path)
    return out_path


def _summary(ws, result, f, title):
    ws.title = "Summary"
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Run date {date.today().isoformat()}, path: {result.path}"])
    ws.append([])
    ws.append(["Object", "Records", "With findings", "Readiness %"])
    for c in ws[4]:
        c.font = HEAD
    for r in result.readiness.to_dict("records"):
        ws.append([r["object"], r["total"], r["records_with_findings"], r["readiness"]])
    ws.append([])
    start = ws.max_row + 1
    ws.append(["Rule", "Title", "blocker", "error", "warning", "info", "Total"])
    for c in ws[start]:
        c.font = HEAD
    counts = f.pivot_table(index="rule", columns="severity", values="key", aggfunc="count", fill_value=0)
    for rid in sorted(REGISTRY):
        if rid in counts.index:
            vals = [int(counts.loc[rid].get(s, 0)) for s in ("blocker", "error", "warning", "info")]
            ws.append([rid, REGISTRY[rid].title, *vals, sum(vals)])
    for col, w in zip("ABCDEFG", (14, 42, 10, 10, 10, 10, 10)):
        ws.column_dimensions[col].width = w


def _worklist(ws, f):
    work = f.rename(columns={"key": "number"})
    for col in ("status", "owner", "comment"):
        if col not in work:
            work[col] = ""  # filled in by the departments, or carried over from the web app
    _sheet(ws, work, {"message": 70, "fix": 55, "comment": 40})
    sev_col = list(work.columns).index("severity") + 1
    for row in ws.iter_rows(min_row=2, min_col=sev_col, max_col=sev_col):
        if row[0].value in SEV_FILL:
            row[0].fill = SEV_FILL[row[0].value]


def _duplicates(ws, f, tables):
    _sheet(ws, duplicate_pairs(f, tables))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top")


def _rules(ws):
    rules = pd.DataFrame([{"rule": r.id, "title": r.title, "applies to": " / ".join(r.paths),
                           "why": " ".join(r.description.split())} for r in REGISTRY.values()])
    _sheet(ws, rules, {"why": 100})
    for row in ws.iter_rows(min_row=2):
        row[-1].alignment = Alignment(wrap_text=True, vertical="top")


def duplicate_pairs(findings, tables):
    cols = ["NAME1", "NAME2", "STRAS", "PSTLZ", "ORT01", "STCEG"]
    lookup = {
        "customer": tables["kna1"].drop_duplicates("KUNNR").set_index("KUNNR"),
        "vendor": tables["lfa1"].drop_duplicates("LIFNR").set_index("LIFNR"),
    }
    rows = []
    d = findings[findings["rule"] == "R01"]
    for r in d.to_dict("records"):
        m = lookup[r["object"]]
        a, b = m.loc[r["related"]], m.loc[r["key"]]
        score = float(r["message"].split("score ")[1].split(":")[0]) if "score " in r["message"] else None
        row = {"object": r["object"], "score": score, "decision": "",
               "number A": r["related"], "number B": r["key"]}
        for c in cols:
            row[f"{c} A"] = a[c]
            row[f"{c} B"] = b[c]
        rows.append(row)
    return pd.DataFrame(rows)
